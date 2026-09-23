"""Persistent Excel reporting for insurance quote comparisons."""

import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook  # pylint: disable=import-error
from openpyxl.styles import Alignment, Font, PatternFill  # pylint: disable=import-error

QUOTE_WORKBOOK_PATH = Path("insurance_quotes.xlsx")
SHEET_NAME = "Quotes"
HEADERS = (
    "Provider",
    "Quote Timestamp",
    "Status",
    "First Name",
    "Last Name",
    "Vehicle Registration",
    "Quote Reference",
    "Cover Type",
    "Payment Schedule",
    "Price (EUR)",
    "Details",
    "Error",
)


@dataclass(frozen=True)
class QuoteRun:
    """Metadata shared by every quote row from one provider run."""

    quote_reference: str = ""
    status: str = "Success"
    error: str = ""
    timestamp: Optional[datetime] = None


def _price_as_number(value):
    """Convert a displayed euro price to a numeric workbook value."""
    if isinstance(value, (int, float)):
        return float(value)
    if not value:
        return None
    match = re.search(r"\d[\d,]*(?:\.\d{1,2})?", str(value))
    if not match:
        return None
    return float(match.group(0).replace(",", ""))


def _prepare_sheet(workbook):
    """Create or validate the quote sheet and apply its stable presentation."""
    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
    else:
        sheet = workbook.create_sheet(SHEET_NAME, 0)

    if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
        for column, header in enumerate(HEADERS, start=1):
            sheet.cell(1, column).value = header
    else:
        actual_headers = tuple(sheet.cell(1, column).value for column in range(1, 13))
        if actual_headers != HEADERS:
            raise ValueError(
                "The Quotes worksheet has unexpected columns; expected: "
                + ", ".join(HEADERS)
            )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False
    widths = (22, 20, 12, 16, 18, 22, 18, 30, 20, 15, 48, 48)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    return sheet


def _save_atomically(workbook, workbook_path):
    """Save without risking a partially written quote workbook."""
    workbook_path = Path(workbook_path)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=workbook_path.parent,
        prefix=f".{workbook_path.stem}-",
        suffix=".xlsx",
        delete=False,
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, workbook_path)
        os.chmod(workbook_path, 0o644)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def ensure_quote_workbook(workbook_path=QUOTE_WORKBOOK_PATH):
    """Create the persistent workbook once without clearing existing quotes."""
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path) if workbook_path.exists() else Workbook()
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) == 1:
        workbook.remove(workbook["Sheet"])
    _prepare_sheet(workbook)
    _save_atomically(workbook, workbook_path)


def upsert_provider_quotes(
    provider,
    data,
    quotes,
    run=None,
    workbook_path=QUOTE_WORKBOOK_PATH,
):
    """Replace one provider's rows while preserving every other provider."""
    run = run or QuoteRun()
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path) if workbook_path.exists() else Workbook()
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) == 1:
        workbook.remove(workbook["Sheet"])
    sheet = _prepare_sheet(workbook)

    for row_number in range(sheet.max_row, 1, -1):
        if sheet.cell(row_number, 1).value == provider:
            sheet.delete_rows(row_number)

    quote_time = run.timestamp or datetime.now()
    rows = quotes or [{}]
    for quote in rows:
        details = quote.get("details", "")
        if isinstance(details, (list, tuple)):
            details = " | ".join(str(detail) for detail in details if detail)
        sheet.append(
            (
                provider,
                quote_time,
                run.status,
                data.get("first_name", ""),
                data.get("last_name", ""),
                data.get("car_registration", ""),
                run.quote_reference,
                quote.get("cover", ""),
                quote.get("payment", ""),
                _price_as_number(quote.get("price")),
                details,
                str(run.error) if run.error else "",
            )
        )
        row_number = sheet.max_row
        sheet.cell(row_number, 2).number_format = "dd/mm/yyyy hh:mm"
        sheet.cell(row_number, 10).number_format = "€#,##0.00"
        for cell in sheet[row_number]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top")
        sheet.cell(row_number, 11).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row_number, 12).alignment = Alignment(vertical="top", wrap_text=True)

    sheet.auto_filter.ref = f"A1:L{sheet.max_row}"
    _save_atomically(workbook, workbook_path)


def record_provider_failure(provider, data, error, workbook_path=QUOTE_WORKBOOK_PATH):
    """Replace the attempted provider with a failure record."""
    upsert_provider_quotes(
        provider,
        data,
        quotes=[],
        run=QuoteRun(status="Failed", error=str(error)),
        workbook_path=workbook_path,
    )
