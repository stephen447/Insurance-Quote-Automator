"""Tests for persistent provider-scoped Excel quote updates."""

import shutil
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook  # pylint: disable=import-error

from helper_functions.excel_report import (
    record_provider_failure,
    upsert_provider_quotes,
)


class ExcelReportTests(unittest.TestCase):
    """Verify reruns replace only the provider that was run."""

    def setUp(self):
        self.temporary_directory = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.temporary_directory, ignore_errors=True)
        self.workbook_path = self.temporary_directory / "quotes.xlsx"
        self.data = {
            "first_name": "Test",
            "last_name": "Driver",
            "car_registration": "12D12345",
        }

    def _rows(self):
        """Read the generated data rows and close the workbook immediately."""
        with closing(load_workbook(self.workbook_path, data_only=True)) as workbook:
            sheet = workbook["Quotes"]
            return list(sheet.iter_rows(min_row=2, values_only=True))

    def test_upsert_replaces_only_matching_provider(self):
        """Rerunning Allianz must preserve AXA's latest rows."""
        upsert_provider_quotes(
            "Allianz Insurance",
            self.data,
            [{"cover": "Comprehensive", "payment": "Annual", "price": "€500"}],
            workbook_path=self.workbook_path,
        )
        upsert_provider_quotes(
            "AXA Insurance",
            self.data,
            [{"cover": "Comprehensive", "payment": "Pay in full", "price": "€600"}],
            workbook_path=self.workbook_path,
        )
        upsert_provider_quotes(
            "Allianz Insurance",
            self.data,
            [{"cover": "Comprehensive", "payment": "Annual", "price": "€450"}],
            workbook_path=self.workbook_path,
        )

        rows = self._rows()
        self.assertEqual(len(rows), 2)
        by_provider = {row[0]: row for row in rows}
        self.assertEqual(by_provider["Allianz Insurance"][9], 450)
        self.assertEqual(by_provider["AXA Insurance"][9], 600)

    def test_failure_replaces_only_attempted_provider(self):
        """A failed AXA rerun must preserve Allianz's successful rows."""
        upsert_provider_quotes(
            "AXA Insurance",
            self.data,
            [{"cover": "Comprehensive", "price": "€600"}],
            workbook_path=self.workbook_path,
        )
        upsert_provider_quotes(
            "Allianz Insurance",
            self.data,
            [{"cover": "Comprehensive", "price": "€500"}],
            workbook_path=self.workbook_path,
        )
        record_provider_failure(
            "AXA Insurance",
            self.data,
            "Quote form changed",
            workbook_path=self.workbook_path,
        )

        rows = self._rows()
        by_provider = {row[0]: row for row in rows}
        self.assertEqual(len(rows), 2)
        self.assertEqual(by_provider["AXA Insurance"][2], "Failed")
        self.assertEqual(by_provider["AXA Insurance"][11], "Quote form changed")
        self.assertEqual(by_provider["Allianz Insurance"][9], 500)


if __name__ == "__main__":
    unittest.main()
